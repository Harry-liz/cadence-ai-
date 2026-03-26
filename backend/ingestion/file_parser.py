from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class ParseResult:
    raw_text: str = ""
    structured_rows: list[dict] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    file_type: str = "unknown"
    metadata: dict = field(default_factory=dict)


EXTENSION_MAP: dict[str, str] = {
    ".xlsx": "xlsx",
    ".xls": "xlsx",
    ".csv": "csv",
    ".pdf": "pdf",
    ".docx": "docx",
    ".doc": "docx",
    ".pptx": "pptx",
    ".txt": "txt",
    ".jpg": "image",
    ".jpeg": "image",
    ".png": "image",
    ".webp": "image",
    ".bmp": "image",
}


def detect_file_type(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    return EXTENSION_MAP.get(ext, "unknown")


def parse_file(file_path: str, file_type: str | None = None) -> ParseResult:
    """Parse a file into a unified ParseResult."""
    if not file_type:
        file_type = detect_file_type(file_path)

    parser = _PARSERS.get(file_type)
    if not parser:
        return ParseResult(
            file_type=file_type,
            metadata={"error": f"Unsupported file type: {file_type}"},
        )
    return parser(file_path)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _parse_excel(file_path: str) -> ParseResult:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ParseResult(file_type="xlsx", metadata={"error": "openpyxl not installed"})

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return ParseResult(file_type="xlsx", metadata={"error": "No active sheet"})

    rows_iter = ws.iter_rows(values_only=True)
    raw_headers = next(rows_iter, None)
    if not raw_headers:
        return ParseResult(file_type="xlsx", metadata={"error": "Empty sheet"})

    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(raw_headers)]

    structured_rows: list[dict] = []
    for row_values in rows_iter:
        if all(v is None for v in row_values):
            continue
        row_dict = {}
        for header, value in zip(headers, row_values):
            row_dict[header] = str(value).strip() if value is not None else ""
        structured_rows.append(row_dict)

    wb.close()
    return ParseResult(
        structured_rows=structured_rows,
        headers=headers,
        file_type="xlsx",
        metadata={
            "sheet_name": ws.title,
            "row_count": len(structured_rows),
            "column_count": len(headers),
        },
    )


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def _parse_csv(file_path: str) -> ParseResult:
    encodings = ["utf-8-sig", "utf-8", "gbk", "gb2312", "latin-1"]
    content = None
    used_encoding = None

    for enc in encodings:
        try:
            with open(file_path, encoding=enc) as f:
                content = f.read()
            used_encoding = enc
            break
        except (UnicodeDecodeError, LookupError):
            continue

    if content is None:
        return ParseResult(file_type="csv", metadata={"error": "Unable to decode file"})

    reader = csv.DictReader(io.StringIO(content))
    headers = list(reader.fieldnames or [])
    structured_rows = list(reader)

    return ParseResult(
        structured_rows=structured_rows,
        headers=headers,
        file_type="csv",
        metadata={
            "encoding": used_encoding,
            "row_count": len(structured_rows),
            "column_count": len(headers),
        },
    )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _parse_pdf(file_path: str) -> ParseResult:
    try:
        import pdfplumber
    except ImportError:
        return ParseResult(file_type="pdf", metadata={"error": "pdfplumber not installed"})

    all_text: list[str] = []
    table_rows: list[dict] = []
    table_headers: list[str] = []

    with pdfplumber.open(file_path) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                all_text.append(text)

            tables = page.extract_tables()
            for table in tables:
                if not table or len(table) < 2:
                    continue
                headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(table[0])]
                if not table_headers:
                    table_headers = headers
                for row in table[1:]:
                    row_dict = {}
                    for header, value in zip(headers, row):
                        row_dict[header] = str(value).strip() if value else ""
                    table_rows.append(row_dict)

    raw_text = "\n\n".join(all_text)

    if len(raw_text.strip()) < 50 and not table_rows:
        return ParseResult(
            raw_text=raw_text,
            file_type="pdf",
            metadata={
                "page_count": page_count,
                "file_path": file_path,
                "needs_multimodal": True,
                "warning": "Very little text extracted — scanned PDF, will use multimodal LLM.",
            },
        )

    return ParseResult(
        raw_text=raw_text,
        structured_rows=table_rows,
        headers=table_headers,
        file_type="pdf",
        metadata={
            "page_count": page_count,
            "text_length": len(raw_text),
            "table_row_count": len(table_rows),
        },
    )


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------

def _parse_docx(file_path: str) -> ParseResult:
    try:
        from docx import Document
    except ImportError:
        return ParseResult(file_type="docx", metadata={"error": "python-docx not installed"})

    doc = Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    raw_text = "\n\n".join(paragraphs)

    table_rows: list[dict] = []
    table_headers: list[str] = []
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if not table_headers:
            table_headers = headers
        for row in table.rows[1:]:
            row_dict = {}
            for header, cell in zip(headers, row.cells):
                row_dict[header] = cell.text.strip()
            table_rows.append(row_dict)

    return ParseResult(
        raw_text=raw_text,
        structured_rows=table_rows,
        headers=table_headers,
        file_type="docx",
        metadata={
            "paragraph_count": len(paragraphs),
            "text_length": len(raw_text),
            "table_row_count": len(table_rows),
        },
    )


# ---------------------------------------------------------------------------
# PowerPoint (.pptx)
# ---------------------------------------------------------------------------

def _parse_pptx(file_path: str) -> ParseResult:
    try:
        from pptx import Presentation
    except ImportError:
        return ParseResult(file_type="pptx", metadata={"error": "python-pptx not installed"})

    prs = Presentation(file_path)
    texts: list[str] = []
    for slide in prs.slides:
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_texts.append(text)
        if slide_texts:
            texts.append("\n".join(slide_texts))

    raw_text = "\n\n".join(texts)
    return ParseResult(
        raw_text=raw_text,
        file_type="pptx",
        metadata={
            "slide_count": len(prs.slides),
            "text_length": len(raw_text),
        },
    )


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------

def _parse_txt(file_path: str) -> ParseResult:
    encodings = ["utf-8-sig", "utf-8", "gbk", "latin-1"]
    for enc in encodings:
        try:
            with open(file_path, encoding=enc) as f:
                content = f.read()
            return ParseResult(
                raw_text=content,
                file_type="txt",
                metadata={"encoding": enc, "text_length": len(content)},
            )
        except (UnicodeDecodeError, LookupError):
            continue
    return ParseResult(file_type="txt", metadata={"error": "Unable to decode file"})


# ---------------------------------------------------------------------------
# Image (OCR) — optional, graceful degradation
# ---------------------------------------------------------------------------

def _parse_image(file_path: str) -> ParseResult:
    raw_text = ""

    try:
        from paddleocr import PaddleOCR
        ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
        result = ocr.ocr(file_path, cls=True)
        lines: list[str] = []
        if result:
            for page in result:
                if page:
                    for line in page:
                        if line and len(line) >= 2:
                            lines.append(line[1][0])
        raw_text = "\n".join(lines)
    except (ImportError, Exception):
        return ParseResult(
            file_type="image",
            metadata={
                "warning": "PaddleOCR not available. Image will be passed to LLM multimodal.",
                "file_path": file_path,
                "needs_multimodal": True,
            },
        )

    return ParseResult(
        raw_text=raw_text,
        file_type="image",
        metadata={
            "text_length": len(raw_text),
            "ocr_engine": "paddleocr",
        },
    )


# ---------------------------------------------------------------------------
# Parser registry
# ---------------------------------------------------------------------------

_PARSERS: dict[str, callable] = {
    "xlsx": _parse_excel,
    "csv": _parse_csv,
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "pptx": _parse_pptx,
    "txt": _parse_txt,
    "image": _parse_image,
}
